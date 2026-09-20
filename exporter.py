"""发票数据导出工具 - Excel/CSV 写入"""

import csv
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    '文件名', '发票号码', '发票代码', '开票日期', '销售方', '销售方税号',
    '购买方', '购买方税号', '商品名称', '规格型号', '单位', '数量',
    '单价', '金额', '税率', '税额', '价税合计', '车辆识别代号/车架号码'
]

COLUMN_WIDTHS = {
    '文件名': 30, '发票号码': 25, '发票代码': 15, '开票日期': 15,
    '销售方': 35, '销售方税号': 20, '购买方': 35, '购买方税号': 20,
    '商品名称': 30, '规格型号': 20, '单位': 10, '数量': 10,
    '单价': 18, '金额': 15, '税率': 10, '税额': 15,
    '价税合计': 15, '车辆识别代号/车架号码': 25,
}

NUMBER_COLS = ['税额', '价税合计', '单价', '金额', '数量']


def write_to_csv(data, filepath):
    """将数据写入CSV文件"""
    with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        writer.writerow(HEADERS)
        for row in data:
            writer.writerow([row.get(h, '') for h in HEADERS])


def write_to_excel(data, filepath):
    """将数据写入Excel文件"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '发票信息'
    ws.append(HEADERS)

    # 表头样式
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='CCE5FF', end_color='CCE5FF', fill_type='solid')

    # 列宽
    for col, header in enumerate(HEADERS, 1):
        if header in COLUMN_WIDTHS:
            ws.column_dimensions[get_column_letter(col)].width = COLUMN_WIDTHS[header]

    # 写入数据
    for row_data in data:
        ws.append([row_data.get(h, '') for h in HEADERS])

    # 格式化数字列
    col_indices = {h: i + 1 for i, h in enumerate(HEADERS)}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for col_name in NUMBER_COLS:
            if col_name in col_indices:
                cell = row[col_indices[col_name] - 1]
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00'

    wb.save(filepath)
    return filepath


def write_data(data, filepath, fmt='xlsx'):
    """统一导出入口:fmt 为 'xlsx' 或 'csv'"""
    if fmt == 'csv':
        write_to_csv(data, filepath)
    else:
        write_to_excel(data, filepath)
    return filepath
