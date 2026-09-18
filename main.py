#!/usr/bin/env python3
"""发票信息批量提取工具"""

import os
import argparse
import csv
import openpyxl
from datetime import datetime
from invoice_extractor import InvoiceExtractor


def main():
    parser = argparse.ArgumentParser(description='发票信息批量提取工具')
    parser.add_argument('input_path', help='PDF文件路径或包含PDF文件的文件夹路径')
    parser.add_argument('--format', '-f', choices=['csv', 'xlsx'], default='xlsx',
                        help='输出格式 (默认: xlsx)')
    parser.add_argument('--output', '-o', help='输出文件路径')
    args = parser.parse_args()

    # 判断是文件还是目录
    if os.path.isfile(args.input_path) and args.input_path.lower().endswith('.pdf'):
        pdf_files = [os.path.basename(args.input_path)]
        input_dir = os.path.dirname(args.input_path) or os.getcwd()
    elif os.path.isdir(args.input_path):
        pdf_files = [f for f in os.listdir(args.input_path) if f.lower().endswith('.pdf')]
        input_dir = args.input_path
    else:
        print("错误：输入路径必须是PDF文件或包含PDF文件的文件夹")
        return

    if not pdf_files:
        print("错误：未找到任何PDF文件")
        return

    output_dir = os.path.dirname(os.path.abspath(args.output)) if args.output else os.path.dirname(os.path.abspath(__file__))
    output_filename = os.path.basename(args.output) if args.output else None

    results = []
    success_count = 0
    fail_count = 0

    print(f"找到 {len(pdf_files)} 个PDF文件")
    print("-" * 50)

    for pdf_file in pdf_files:
        pdf_path = os.path.join(input_dir, pdf_file)
        print(f"正在处理: {pdf_file}")

        try:
            extractor = InvoiceExtractor(pdf_path)
            rows = extractor.extract_all_rows()

            # 填充文件名
            for row in rows:
                row['文件名'] = pdf_file
                results.append(row)

            success_count += 1
            print(f"  ✓ 提取成功 ({len(rows)}行)")
        except Exception as e:
            print(f"  ✗ 处理失败: {str(e)}")
            fail_count += 1
            results.append({
                '文件名': pdf_file,
                '发票号码': '提取失败',
                '发票代码': '提取失败',
                '开票日期': '提取失败',
                '销售方': '提取失败',
                '销售方税号': '提取失败',
                '购买方': '提取失败',
                '购买方税号': '提取失败',
                '商品名称': '提取失败',
                '规格型号': '提取失败',
                '单位': '提取失败',
                '数量': '提取失败',
                '单价': '提取失败',
                '金额': '提取失败',
                '税率': '提取失败',
                '税额': '提取失败',
                '价税合计': '提取失败',
                '车辆识别代号/车架号码': '提取失败',
            })

    print("-" * 50)

    # 确定输出路径
    if output_filename:
        output_path = os.path.join(output_dir, output_filename)
    else:
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        if args.format == 'xlsx':
            output_path = os.path.join(output_dir, f'发票信息_{timestamp}.xlsx')
        else:
            output_path = os.path.join(output_dir, f'发票信息_{timestamp}.csv')

    # 定义表头
    headers = ['文件名', '发票号码', '发票代码', '开票日期', '销售方', '销售方税号',
               '购买方', '购买方税号', '商品名称', '规格型号', '单位', '数量',
               '单价', '金额', '税率', '税额', '价税合计', '车辆识别代号/车架号码']

    if args.format == 'xlsx':
        write_to_excel(results, output_path, headers)
    else:
        write_to_csv(results, output_path, headers)

    print(f"处理完成！成功: {success_count}个文件, 失败: {fail_count}个, 共{len(results)}行数据")
    print(f"结果已保存到: {output_path}")


def write_to_csv(data, filepath, headers):
    """将数据写入CSV文件"""
    with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for row in data:
            writer.writerow([row.get(h, '') for h in headers])


def write_to_excel(data, filepath, headers):
    """将数据写入Excel文件"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '发票信息'

    ws.append(headers)

    # 设置表头样式
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = openpyxl.styles.Font(bold=True)
        cell.fill = openpyxl.styles.PatternFill(start_color='CCE5FF', end_color='CCE5FF', fill_type='solid')

    # 设置列宽
    column_widths = {
        '文件名': 30,
        '发票号码': 25,
        '发票代码': 15,
        '开票日期': 15,
        '销售方': 35,
        '销售方税号': 20,
        '购买方': 35,
        '购买方税号': 20,
        '商品名称': 30,
        '规格型号': 20,
        '单位': 10,
        '数量': 10,
        '单价': 18,
        '金额': 15,
        '税率': 10,
        '税额': 15,
        '价税合计': 15,
        '车辆识别代号/车架号码': 25,
    }

    for col, header in enumerate(headers, 1):
        if header in column_widths:
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = column_widths[header]

    # 写入数据
    for row_data in data:
        ws.append([row_data.get(h, '') for h in headers])

    # 格式化数字列
    number_cols = ['税额', '价税合计', '单价', '金额', '数量']
    col_indices = {h: i + 1 for i, h in enumerate(headers)}

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for col_name in number_cols:
            if col_name in col_indices:
                cell = row[col_indices[col_name] - 1]
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00'

    wb.save(filepath)


if __name__ == '__main__':
    main()
